#!/usr/bin/env python3
"""Generate the final clinical interpretation DOCX from an hgene run.

Missing or malformed clinical metadata are intentionally non-blocking: the
corresponding placeholders are replaced with empty strings. Only an unusable
Word template or an unwritable output file is fatal.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator, Mapping, Sequence

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook


PLACEHOLDER_RE = re.compile(r"_[A-Z0-9]+(?:-[A-Z0-9]+)*_")
RESISTANCE_MARKER_RE = re.compile(
    r"^_(?P<kind>MUT|INT)-(?P<drug>[A-Z0-9]+)-(?P<gene>[A-Z0-9]+)_$"
)

DRUG_LABELS = {
    "GCV": "ganciclovir",
    "CDV": "cidofovir",
    "FOS": "foscarnet",
    "MBV": "maribavir",
    "LMV": "letermovir",
    "ACV": "aciclovir",
    "VACV": "valaciclovir",
    "PCV": "penciclovir",
    "FCV": "famciclovir",
    "BVDU": "brivudine",
}

DRUG_ALIASES = {
    "GCV": ("gcv", "ganciclovir"),
    "CDV": ("cdv", "cidofovir"),
    "FOS": ("fos", "foscarnet", "phosphonoformate"),
    "MBV": ("mbv", "maribavir"),
    "LMV": ("lmv", "letermovir"),
    "ACV": ("acv", "acyclovir", "aciclovir"),
    "VACV": ("vacv", "valacyclovir", "valaciclovir"),
    "PCV": ("pcv", "penciclovir"),
    "FCV": ("fcv", "famciclovir"),
    "BVDU": ("bvdu", "brivudine"),
}

DRUG_PREPOSITIONS = {
    "ACV": "à l’aciclovir",
}

DRUG_ORDER = ("GCV", "CDV", "FOS", "MBV", "LMV", "ACV", "VACV", "PCV", "FCV", "BVDU")


@dataclass(frozen=True)
class Variant:
    chrom: str
    pos: int
    ref: str
    alt: str
    gene: str
    aa_raw: str
    aa_report: str
    dna_raw: str
    dna_report: str
    consequence: str
    af: float | None


@dataclass(frozen=True)
class ResistanceMatch:
    variant: Variant
    drug: str
    mutation_type: str


@dataclass(frozen=True)
class MutationTypeInfo:
    aggregate_state: str
    aggregate_rank: int
    table_state: str | None = None
    table_qualifier: str = ""
    aggregate_qualifier: str = ""
    recognized: bool = True


MUTATION_TYPE_RULES: Mapping[str, MutationTypeInfo] = {
    "Natural polymorphism": MutationTypeInfo(
        "susceptible", 0, table_state="natural_polymorphism"
    ),
    "No resistance": MutationTypeInfo("susceptible", 0),
    "Low level resistance": MutationTypeInfo(
        "resistant",
        1,
        table_qualifier="faible niveau",
        aggregate_qualifier="faible niveau",
    ),
    "Intermediate level resistance": MutationTypeInfo(
        "resistant",
        2,
        table_qualifier="niveau intermédiaire",
        aggregate_qualifier="niveau intermédiaire",
    ),
    "High level resistance": MutationTypeInfo(
        "resistant",
        3,
        table_qualifier="haut niveau",
        aggregate_qualifier="haut niveau",
    ),
    "Non-viable": MutationTypeInfo("non_viable", 0),
    "Possible hypersensitivity, no resistance": MutationTypeInfo(
        "susceptible",
        0,
        table_qualifier="hypersensibilité possible",
    ),
    "No more than Low level resistance": MutationTypeInfo(
        "resistant",
        1,
        table_qualifier="niveau faible au maximum",
        aggregate_qualifier="faible niveau",
    ),
    "At least Intermediate level resistance": MutationTypeInfo(
        "resistant",
        2,
        table_qualifier="niveau au moins intermédiaire",
    ),
    "At least High level resistance": MutationTypeInfo(
        "resistant",
        3,
        table_qualifier="niveau au moins élevé",
        aggregate_qualifier="haut niveau",
    ),
    "Range of values of No resistance": MutationTypeInfo(
        "susceptible",
        0,
        table_qualifier="plage de valeurs sans résistance",
    ),
    "Range of values of High level resistance": MutationTypeInfo(
        "resistant",
        3,
        table_qualifier="plage de valeurs de haut niveau",
        aggregate_qualifier="haut niveau",
    ),
    "Range of values between No resistance and Intermediate level resistance": (
        MutationTypeInfo("indeterminate", 2, table_state="variable")
    ),
    "Range of values between Low level resistance and Intermediate level resistance": (
        MutationTypeInfo(
            "resistant",
            2,
            table_qualifier="niveau faible à intermédiaire",
            aggregate_qualifier="niveau intermédiaire",
        )
    ),
    "Range of values between Intermediate level resistance and High level resistance": (
        MutationTypeInfo(
            "resistant",
            3,
            table_qualifier="niveau intermédiaire à élevé",
            aggregate_qualifier="haut niveau",
        )
    ),
}


@dataclass(frozen=True)
class CoverageStatus:
    low20: bool
    low100: bool


def warn(message: str) -> None:
    print(f"WARNING: {message}", file=sys.stderr)


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.casefold() in {"", ".", "na", "n/a", "nan", "none", "unknown"}:
        return ""
    return text


def first_nonempty(*values: object) -> str:
    for value in values:
        text = clean(value)
        if text:
            return text
    return ""


def sha256_file(path: Path | None) -> str:
    if path is None or not path.is_file():
        return ""
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)
    except OSError as exc:
        warn(f"cannot calculate SHA-256 for {path}: {exc}")
        return ""
    return digest.hexdigest()


def open_text(path: Path):
    if path.suffix == ".gz":
        return gzip.open(path, "rt", encoding="utf-8", errors="replace")
    return path.open("rt", encoding="utf-8", errors="replace")


def normalize_gene(value: object) -> str:
    text = clean(value).upper()
    match = re.search(r"\b(UL\d+[A-Z]?)", text)
    if match:
        return match.group(1)
    return re.sub(r"[^A-Z0-9]+", "", text)


def report_aa_change(chrom: str, aa_change: str) -> str:
    aa_change = clean(aa_change)
    if chrom not in {"UL89_1-HHV5", "UL89_3-HHV5"}:
        return aa_change
    match = re.fullmatch(r"(\d+)([A-Za-z*]+)>(\d+)([A-Za-z*]+)", aa_change)
    if not match:
        return aa_change
    return (
        f"{int(match.group(1)) + 296}{match.group(2)}>"
        f"{int(match.group(3)) + 296}{match.group(4)}"
    )


def report_dna_change(chrom: str, pos: int, ref: str, alt: str, raw: str) -> str:
    if chrom in {"UL89_1-HHV5", "UL89_3-HHV5"}:
        return f"{pos + 888}{ref}>{alt}"
    return first_nonempty(raw, f"{pos}{ref}>{alt}")


def parse_info_field(raw: str) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in raw.split(";"):
        if "=" in item:
            key, value = item.split("=", 1)
            result[key] = value
    return result


def parse_float(value: object) -> float | None:
    try:
        return float(str(value).split(",", 1)[0])
    except (TypeError, ValueError):
        return None


def read_vcf(path: Path | None) -> tuple[list[Variant], dict[str, str], bool]:
    if path is None or not path.is_file():
        warn(f"VCF not found: {path}")
        return [], {}, False

    variants: list[Variant] = []
    metadata: dict[str, str] = {}
    try:
        with open_text(path) as handle:
            for line in handle:
                line = line.rstrip("\n")
                if line.startswith("##") and "=" in line:
                    key, value = line[2:].split("=", 1)
                    metadata[key] = value
                    continue
                if not line or line.startswith("#"):
                    continue
                fields = line.split("\t")
                if len(fields) < 8:
                    warn(f"ignoring malformed VCF line in {path}")
                    continue
                chrom, pos_raw, _vid, ref, alt = fields[:5]
                try:
                    pos = int(pos_raw)
                except ValueError:
                    warn(f"ignoring VCF line with invalid position: {pos_raw}")
                    continue
                info = parse_info_field(fields[7])
                af = parse_float(info.get("AF"))
                bcsq_values = clean(info.get("BCSQ")).split(",")
                for bcsq in bcsq_values:
                    if not bcsq or bcsq.startswith("@"):
                        continue
                    parts = bcsq.split("|")
                    if len(parts) < 7:
                        warn(f"ignoring malformed BCSQ annotation at {chrom}:{pos}")
                        continue
                    consequence = parts[0]
                    gene_raw = parts[1]
                    aa_raw = parts[5]
                    dna_raw = parts[6]
                    variants.append(
                        Variant(
                            chrom=chrom,
                            pos=pos,
                            ref=ref,
                            alt=alt.split(",", 1)[0],
                            gene=normalize_gene(gene_raw or chrom),
                            aa_raw=clean(aa_raw),
                            aa_report=report_aa_change(chrom, aa_raw),
                            dna_raw=clean(dna_raw),
                            dna_report=report_dna_change(
                                chrom, pos, ref, alt.split(",", 1)[0], dna_raw
                            ),
                            consequence=clean(consequence),
                            af=af,
                        )
                    )
    except (OSError, EOFError, UnicodeError) as exc:
        warn(f"cannot read VCF {path}: {exc}")
        return [], metadata, False
    return variants, metadata, True


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", clean(value))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()


def read_resistance_db(
    path: Path | None,
) -> tuple[list[dict[str, str]], dict[str, str], bool]:
    if path is None or not path.is_file():
        warn(f"resistance database not found: {path}")
        return [], {}, False
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        data_sheet_name = next(
            (name for name in workbook.sheetnames if name.casefold() != "metadata"),
            workbook.sheetnames[0],
        )
        sheet = workbook[data_sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [normalize_header(value) for value in next(rows)]
        records: list[dict[str, str]] = []
        for row in rows:
            record = {
                header: clean(value)
                for header, value in zip(headers, row)
                if header
            }
            if any(record.get(key) for key in ("drug", "aa_chg_bcsq", "dna_chg_bcsq", "aa_change")):
                records.append(record)

        metadata: dict[str, str] = {}
        if "metadata" in workbook.sheetnames:
            meta_rows = workbook["metadata"].iter_rows(values_only=True)
            next(meta_rows, None)
            for row in meta_rows:
                if len(row) >= 2 and clean(row[0]):
                    metadata[clean(row[0])] = clean(row[1])
        workbook.close()
    except Exception as exc:  # openpyxl exposes several format-specific exceptions
        warn(f"cannot read resistance database {path}: {exc}")
        return [], {}, False

    usable = bool(records)
    if not usable:
        warn(f"no usable resistance records found in {path}")
    return records, metadata, usable


def parse_bool(value: object) -> bool:
    return clean(value).casefold() in {"1", "true", "t", "yes", "y", "oui"}


def read_coverage(path: Path | None) -> tuple[dict[str, CoverageStatus], bool]:
    if path is None or not path.is_file():
        warn(f"coverage summary not found: {path}")
        return {}, False
    result: dict[str, CoverageStatus] = {}
    try:
        with path.open("rt", encoding="utf-8-sig", errors="replace", newline="") as handle:
            for row in csv.DictReader(handle, delimiter="\t"):
                gene = normalize_gene(row.get("gene"))
                if not gene:
                    continue
                low20 = parse_bool(row.get("low_coverage_20x"))
                low100 = parse_bool(row.get("low_coverage_100x")) or low20
                previous = result.get(gene)
                if previous:
                    low20 = low20 or previous.low20
                    low100 = low100 or previous.low100
                result[gene] = CoverageStatus(low20=low20, low100=low100)
    except (OSError, csv.Error) as exc:
        warn(f"cannot read coverage summary {path}: {exc}")
        return {}, False
    return result, bool(result)


def parse_drugs(value: object) -> list[str]:
    text = unicodedata.normalize("NFKD", clean(value)).casefold()
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    if not text:
        return []
    found: list[str] = []
    for code, aliases in DRUG_ALIASES.items():
        if any(
            re.search(rf"(?<![a-z0-9]){re.escape(alias)}(?![a-z0-9])", text)
            for alias in aliases
        ):
            found.append(code)
    return found


def same_value(left: object, right: object) -> bool:
    return clean(left).casefold() == clean(right).casefold() and bool(clean(left))


def match_variants(
    variants: Sequence[Variant], records: Sequence[Mapping[str, str]]
) -> list[ResistanceMatch]:
    matches: list[ResistanceMatch] = []
    seen: set[tuple[Variant, str, str]] = set()
    for variant in variants:
        for record in records:
            chrom_match = same_value(record.get("chrom"), variant.chrom)
            aa_match = chrom_match and same_value(record.get("aa_chg_bcsq"), variant.aa_raw)
            dna_match = chrom_match and same_value(record.get("dna_chg_bcsq"), variant.dna_raw)
            gene_match = (
                normalize_gene(record.get("gene")) == variant.gene
                and same_value(record.get("aa_change"), variant.aa_report)
            )
            if not (aa_match or dna_match or gene_match):
                continue
            mutation_type = clean(record.get("mutation_type"))
            for drug in parse_drugs(record.get("drug")):
                key = (variant, drug, mutation_type)
                if key not in seen:
                    seen.add(key)
                    matches.append(
                        ResistanceMatch(
                            variant=variant,
                            drug=drug,
                            mutation_type=mutation_type,
                        )
                    )
    return matches


def mutation_type_info(value: str) -> MutationTypeInfo:
    """Interpret only the exact mutation_type values validated for CHARMD."""
    return MUTATION_TYPE_RULES.get(
        clean(value),
        MutationTypeInfo("indeterminate", 0, recognized=False),
    )


def mutation_aggregate_state(value: str) -> str:
    return mutation_type_info(value).aggregate_state


def aggregate_mutation_types(
    matches: Sequence[ResistanceMatch],
) -> MutationTypeInfo | None:
    if not matches:
        return None
    infos = [mutation_type_info(match.mutation_type) for match in matches]
    state_priority = {
        "susceptible": 1,
        "non_viable": 2,
        "indeterminate": 3,
        "resistant": 4,
    }
    return max(
        infos,
        key=lambda info: (
            state_priority.get(info.aggregate_state, 0),
            info.aggregate_rank,
            bool(info.aggregate_qualifier),
        ),
    )


def mutation_interpretation_text(
    drug: str, matches: Sequence[ResistanceMatch]
) -> str:
    info = aggregate_mutation_types(matches)
    if info is None:
        return ""
    table_state = info.table_state or info.aggregate_state
    if table_state == "resistant":
        text = f"Résistant {drug_preposition(drug)}"
    elif table_state == "susceptible":
        text = f"Susceptible {drug_preposition(drug)}"
    elif table_state == "natural_polymorphism":
        text = "Polymorphisme naturel, non associé à une résistance"
    elif table_state == "non_viable":
        text = "Mutant non viable"
    elif table_state == "variable":
        text = (
            "Interprétation variable pour le "
            f"{DRUG_LABELS.get(drug, drug.lower())} : de susceptible à "
            "résistant de niveau intermédiaire"
        )
    else:
        text = (
            "Interprétation indéterminée pour le "
            f"{DRUG_LABELS.get(drug, drug.lower())}"
        )
    if info.table_qualifier:
        text += f" ({info.table_qualifier})"
    return text


def format_protein_change(raw: str) -> str:
    raw = clean(raw)
    match = re.fullmatch(r"(\d+)([A-Za-z*]+)>(\d+)([A-Za-z*]+)", raw)
    if match and match.group(1) == match.group(3):
        return f"p.{match.group(2)}{match.group(1)}{match.group(4)}"
    match = re.fullmatch(r"(\d+)([A-Za-z*]+)>([A-Za-z*]+)", raw)
    if match:
        return f"p.{match.group(2)}{match.group(1)}{match.group(3)}"
    if re.fullmatch(r"[A-Za-z*]+\d+[A-Za-z*]+", raw):
        return f"p.{raw}"
    return raw


def format_mutation(variant: Variant, with_af: bool = False) -> str:
    change = format_protein_change(variant.aa_report)
    if not change:
        change = variant.dna_report
    text = f"{variant.gene}: {change}" if change else variant.gene
    if with_af and variant.af is not None:
        percent = f"{100 * variant.af:.1f}".replace(".", ",")
        text = f"{text} à {percent} %"
    return text


def drug_preposition(code: str) -> str:
    return DRUG_PREPOSITIONS.get(code, f"au {DRUG_LABELS.get(code, code.lower())}")


def join_french(items: Sequence[str]) -> str:
    items = [item for item in items if item]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return ", ".join(items[:-1]) + " et " + items[-1]


def coverage_limit_text(
    genes: Iterable[str], coverage: Mapping[str, CoverageStatus]
) -> str:
    low20 = [gene for gene in sorted(genes) if coverage.get(gene) and coverage[gene].low20]
    low100 = [
        gene
        for gene in sorted(genes)
        if coverage.get(gene) and coverage[gene].low100 and not coverage[gene].low20
    ]
    parts: list[str] = []
    if low20:
        parts.append(f"couverture partiellement <20x pour {join_french(low20)}")
    if low100:
        parts.append(f"couverture partiellement <100x pour {join_french(low100)}")
    return "; ".join(parts)


def group_matches_by_variant(
    matches: Iterable[ResistanceMatch],
) -> list[list[ResistanceMatch]]:
    groups: list[list[ResistanceMatch]] = []
    indexes: dict[tuple[str, int, str, str, float | None], int] = {}
    for match in matches:
        variant = match.variant
        key = (variant.gene, variant.pos, variant.ref, variant.alt, variant.af)
        if key in indexes:
            groups[indexes[key]].append(match)
        else:
            indexes[key] = len(groups)
            groups.append([match])
    return groups


def unique_matches(matches: Iterable[ResistanceMatch]) -> list[ResistanceMatch]:
    return [group[0] for group in group_matches_by_variant(matches)]


def iter_unique_cells(table) -> Iterator:
    seen: set[object] = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            yield cell
            for nested_table in cell.tables:
                yield from iter_unique_cells(nested_table)


def iter_all_cells(doc: Document) -> Iterator:
    for table in doc.tables:
        yield from iter_unique_cells(table)
    for section in doc.sections:
        for container in (section.header, section.footer):
            for table in container.tables:
                yield from iter_unique_cells(table)


def iter_all_paragraphs(doc: Document) -> Iterator:
    yield from doc.paragraphs
    for cell in iter_all_cells(doc):
        yield from cell.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def discover_placeholders(doc: Document) -> set[str]:
    markers: set[str] = set()
    for paragraph in iter_all_paragraphs(doc):
        markers.update(PLACEHOLDER_RE.findall(paragraph.text))
    return markers


def resistance_pairs(markers: Iterable[str]) -> list[tuple[str, str]]:
    pairs: set[tuple[str, str]] = set()
    for marker in markers:
        match = RESISTANCE_MARKER_RE.fullmatch(marker)
        if match:
            pairs.add((match.group("drug"), normalize_gene(match.group("gene"))))
    return sorted(pairs, key=lambda pair: (drug_sort_key(pair[0]), pair[1]))


def drug_sort_key(code: str) -> tuple[int, str]:
    try:
        return DRUG_ORDER.index(code), code
    except ValueError:
        return len(DRUG_ORDER), code


def shade_marker_cells(doc: Document, markers: Iterable[str], fill: str) -> None:
    wanted = set(markers)
    for cell in iter_all_cells(doc):
        if not wanted.intersection(PLACEHOLDER_RE.findall(cell.text)):
            continue
        properties = cell._tc.get_or_add_tcPr()
        for old in list(properties.findall(qn("w:shd"))):
            properties.remove(old)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)


def replace_in_paragraph(paragraph, replacements: Mapping[str, str]) -> None:
    original = "".join(run.text for run in paragraph.runs)
    if not original:
        return
    updated = original
    for marker, value in replacements.items():
        if marker in updated:
            updated = updated.replace(marker, value)
    if updated == original:
        return
    if paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(updated)


def replace_all(doc: Document, replacements: Mapping[str, str]) -> None:
    for paragraph in iter_all_paragraphs(doc):
        replace_in_paragraph(paragraph, replacements)


def parse_sample_metadata(sample_name: str) -> tuple[str, str]:
    match = re.match(r"^([^_]+)_([^_]+)_([^_]+)_([^_]+)_([^.]*)", sample_name)
    if not match:
        warn(f"cannot extract tube ID and patient initials from sample name: {sample_name}")
        return "", ""
    return clean(match.group(3)), clean(match.group(5))


def parse_set_values(values: Sequence[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for item in values:
        if "=" not in item:
            warn(f"ignoring invalid --set value (expected PLACEHOLDER=VALUE): {item}")
            continue
        marker, value = item.split("=", 1)
        marker = marker.strip()
        if not marker.startswith("_"):
            marker = f"_{marker}"
        if not marker.endswith("_"):
            marker = f"{marker}_"
        result[marker] = value
    return result


def build_conclusion(
    pairs: Sequence[tuple[str, str]],
    matches: Sequence[ResistanceMatch],
    coverage: Mapping[str, CoverageStatus],
    vcf_ok: bool,
    db_ok: bool,
    coverage_ok: bool,
) -> str:
    if not vcf_ok or not db_ok:
        return ""

    genes_by_drug: dict[str, set[str]] = {}
    for drug, gene in pairs:
        genes_by_drug.setdefault(drug, set()).add(gene)

    resistant_sentences: list[str] = []
    susceptible_sentences: list[str] = []
    non_viable_sentences: list[str] = []
    limited_sentences: list[str] = []
    absence_susceptible: list[str] = []
    for drug in sorted(genes_by_drug, key=drug_sort_key):
        drug_matches = [match for match in matches if match.drug == drug]
        resistant = unique_matches(
            match
            for match in drug_matches
            if mutation_aggregate_state(match.mutation_type) == "resistant"
        )
        uncertain = unique_matches(
            match
            for match in drug_matches
            if mutation_aggregate_state(match.mutation_type) == "indeterminate"
        )
        non_viable = unique_matches(
            match
            for match in drug_matches
            if mutation_aggregate_state(match.mutation_type) == "non_viable"
        )
        susceptible = unique_matches(
            match
            for match in drug_matches
            if mutation_aggregate_state(match.mutation_type) == "susceptible"
        )
        required_genes = genes_by_drug[drug]
        statuses = [coverage.get(gene) for gene in required_genes]
        coverage_complete = (
            coverage_ok
            and all(status is not None for status in statuses)
            and all(not status.low100 for status in statuses if status is not None)
        )
        coverage_limited = [
            gene
            for gene in sorted(required_genes)
            if coverage.get(gene) is not None and coverage[gene].low100
        ]

        if resistant:
            overall = aggregate_mutation_types(drug_matches)
            detail = join_french(
                [format_mutation(match.variant, with_af=True) for match in resistant]
            )
            mutation_label = (
                "de la mutation de résistance"
                if len(resistant) == 1
                else "des mutations de résistance"
            )
            resistance_level = (
                f" ({overall.aggregate_qualifier})"
                if overall and overall.aggregate_qualifier
                else ""
            )
            sentence = (
                f"Souche résistante {drug_preposition(drug)}{resistance_level} : "
                f"présence {mutation_label} "
                f"{detail}."
            )
            if uncertain:
                uncertain_detail = join_french(
                    [format_mutation(match.variant, with_af=True) for match in uncertain]
                )
                sentence += (
                    " Une interprétation indéterminée est associée à "
                    f"{uncertain_detail}."
                )
            if non_viable:
                non_viable_detail = join_french(
                    [format_mutation(match.variant, with_af=True) for match in non_viable]
                )
                label = "Mutant non viable" if len(non_viable) == 1 else "Mutants non viables"
                sentence += f" {label} : {non_viable_detail}."
            if coverage_limited:
                limit_text = coverage_limit_text(coverage_limited, coverage)
                limit_text = limit_text[:1].upper() + limit_text[1:]
                sentence += (
                    f" {limit_text} : "
                    "la recherche d’autres mutations est limitée."
                )
            resistant_sentences.append(sentence)
        elif uncertain:
            overall = aggregate_mutation_types(uncertain)
            detail = join_french(
                [format_mutation(match.variant, with_af=True) for match in uncertain]
            )
            qualifier = (
                f" ({overall.aggregate_qualifier})"
                if overall and overall.aggregate_qualifier
                else ""
            )
            sentence = (
                "Interprétation non concluante pour le "
                f"{DRUG_LABELS.get(drug, drug.lower())}{qualifier} : "
                f"présence de {detail}."
            )
            if coverage_limited:
                sentence += f" {coverage_limit_text(coverage_limited, coverage)}."
            if non_viable:
                non_viable_detail = join_french(
                    [format_mutation(match.variant, with_af=True) for match in non_viable]
                )
                label = "Mutant non viable" if len(non_viable) == 1 else "Mutants non viables"
                sentence += f" {label} : {non_viable_detail}."
            limited_sentences.append(sentence)
        elif non_viable:
            detail = join_french(
                [format_mutation(match.variant, with_af=True) for match in non_viable]
            )
            label = "Mutant non viable" if len(non_viable) == 1 else "Mutants non viables"
            sentence = f"{label} : {detail}."
            if coverage_limited:
                sentence += f" {coverage_limit_text(coverage_limited, coverage)}."
            non_viable_sentences.append(sentence)
        elif susceptible and coverage_complete:
            detail = join_french(
                [format_mutation(match.variant, with_af=True) for match in susceptible]
            )
            mutation_label = "de la mutation" if len(susceptible) == 1 else "des mutations"
            susceptible_sentences.append(
                f"Souche sensible {drug_preposition(drug)} : présence {mutation_label} "
                f"{detail}, associée à une susceptibilité {drug_preposition(drug)}."
            )
        elif susceptible and coverage_limited:
            detail = join_french(
                [format_mutation(match.variant, with_af=True) for match in susceptible]
            )
            if len(susceptible) == 1:
                mutation_label = "une mutation associée"
                detected_label = "a été détectée"
            else:
                mutation_label = "des mutations associées"
                detected_label = "ont été détectées"
            limited_sentences.append(
                f"Interprétation non concluante pour le "
                f"{DRUG_LABELS.get(drug, drug.lower())} : {mutation_label} à une "
                f"susceptibilité {drug_preposition(drug)} {detected_label} ({detail}), mais "
                f"{coverage_limit_text(coverage_limited, coverage)}."
            )
        elif coverage_complete:
            absence_susceptible.append(DRUG_LABELS.get(drug, drug.lower()))
        elif coverage_limited:
            limited_sentences.append(
                f"Interprétation non concluante pour le {DRUG_LABELS.get(drug, drug.lower())} : "
                f"{coverage_limit_text(coverage_limited, coverage)}."
            )

    sentences = resistant_sentences + susceptible_sentences + non_viable_sentences
    if absence_susceptible:
        sentences.append(
            "Souche sensible aux antiviraux suivants : "
            f"{join_french(absence_susceptible)}. Absence de mutation de résistance "
            "connue dans les gènes analysés."
        )
    sentences.extend(limited_sentences)
    return "\n\n".join(sentences)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Generate the final hgene clinical interpretation DOCX."
    )
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--vcf", type=Path)
    parser.add_argument("--coverage", type=Path)
    parser.add_argument("--resistance-db", type=Path)
    parser.add_argument("--sample-name", default="")
    parser.add_argument("--date-prelevement", default=os.getenv("HG_DATE_PRELEVEMENT", ""))
    parser.add_argument("--tube-id", default=os.getenv("HG_TUBE_ID", ""))
    parser.add_argument("--patient-name", default=os.getenv("HG_PATIENT_NAME", ""))
    parser.add_argument("--materiel", default=os.getenv("HG_MATERIEL", ""))
    parser.add_argument("--analysis-number", default=os.getenv("HG_ANALYSIS_NUMBER", ""))
    parser.add_argument("--database-name", default=os.getenv("HG_DATABASE_NAME", ""))
    parser.add_argument("--ref-accession", default=os.getenv("HG_REF_ACCESSION", ""))
    parser.add_argument("--software-version", default=os.getenv("HG_SOFTWARE_VERSION", ""))
    parser.add_argument("--container-sha256", default=os.getenv("HG_CONTAINER_SHA256", ""))
    parser.add_argument("--database-sha256", default=os.getenv("HG_DATABASE_SHA256", ""))
    parser.add_argument("--conclusion", default=os.getenv("HG_CLINICAL_CONCLUSION", ""))
    parser.add_argument(
        "--set",
        action="append",
        default=[],
        metavar="PLACEHOLDER=VALUE",
        help="override any template placeholder; may be repeated",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    if not args.template.is_file():
        print(f"ERROR: Word template not found: {args.template}", file=sys.stderr)
        return 2

    try:
        doc = Document(args.template)
    except Exception as exc:
        print(f"ERROR: cannot open Word template {args.template}: {exc}", file=sys.stderr)
        return 2

    markers = discover_placeholders(doc)
    pairs = resistance_pairs(markers)

    variants, vcf_meta, vcf_ok = read_vcf(args.vcf)
    db_records, db_meta, db_ok = read_resistance_db(args.resistance_db)
    coverage, coverage_ok = read_coverage(args.coverage)
    matches = match_variants(variants, db_records) if db_ok else []
    unexpected_mutation_types = sorted(
        {
            clean(match.mutation_type) or "<vide>"
            for match in matches
            if not mutation_type_info(match.mutation_type).recognized
        }
    )
    for mutation_type in unexpected_mutation_types:
        warn(
            f"unrecognized mutation_type {mutation_type!r}; "
            "clinical interpretation set to indeterminate"
        )

    sample_name = first_nonempty(
        args.sample_name,
        args.vcf.name.removesuffix(".vcf.gz") if args.vcf else "",
    )
    parsed_tube, parsed_patient = parse_sample_metadata(sample_name) if sample_name else ("", "")

    replacements: dict[str, str] = {marker: "" for marker in markers}
    replacements.update(
        {
            "_DATEPREL_": clean(args.date_prelevement),
            "_IDTUBE_": first_nonempty(args.tube_id, parsed_tube),
            "_IDPATIENT_": first_nonempty(args.patient_name, parsed_patient),
            "_MATERIEL_": clean(args.materiel),
            "_DATABASE_": first_nonempty(args.database_name, db_meta.get("db_name")),
            "_NUMANALYSE_": clean(args.analysis_number),
            "_ANALYSISNUMBER_": clean(args.analysis_number),
            "_ACCESSIONNUMBER_": first_nonempty(
                args.ref_accession, vcf_meta.get("hgene_ref_accession")
            ),
            "_VERSIONLOGICIEL_": first_nonempty(
                args.software_version, vcf_meta.get("hgene_version")
            ),
            "_EMPREINTE_": first_nonempty(
                args.container_sha256, vcf_meta.get("hgene_container_sha256")
            ),
            "_DBEMPREINTE_": first_nonempty(
                args.database_sha256, sha256_file(args.resistance_db)
            ),
        }
    )

    for drug, gene in pairs:
        raw_pair_matches = [
            match
            for match in matches
            if match.drug == drug and match.variant.gene == gene
        ]
        pair_match_groups = group_matches_by_variant(raw_pair_matches)
        pair_matches = [group[0] for group in pair_match_groups]
        mutation_marker = f"_MUT-{drug}-{gene}_"
        interpretation_marker = f"_INT-{drug}-{gene}_"
        mutation_text = "\n".join(
            format_mutation(match.variant) for match in pair_matches
        )
        interpretation_text = "\n".join(
            mutation_interpretation_text(drug, group)
            for group in pair_match_groups
        )
        pair_states = {
            mutation_aggregate_state(match.mutation_type)
            for match in raw_pair_matches
        }
        if pair_states.intersection({"indeterminate", "non_viable"}):
            shade_marker_cells(doc, (mutation_marker, interpretation_marker), "FFF2CC")

        coverage_status = coverage.get(gene)
        if coverage_status and coverage_status.low20:
            shade_marker_cells(doc, (mutation_marker, interpretation_marker), "F4CCCC")
            if mutation_text:
                interpretation_text += (
                    "\nCouverture <20x : recherche d’autres mutations limitée"
                )
            else:
                mutation_text = "Couverture partiellement <20x"
                interpretation_text = "Non interprétable"
        elif coverage_status and coverage_status.low100:
            shade_marker_cells(doc, (mutation_marker, interpretation_marker), "FFF2CC")
            if mutation_text:
                interpretation_text += (
                    "\nCouverture <100x : recherche des variants minoritaires limitée"
                )
            else:
                mutation_text = "Couverture partiellement <100x"
                interpretation_text = "Interprétation limitée"

        replacements[mutation_marker] = mutation_text
        replacements[interpretation_marker] = interpretation_text

    generated_conclusion = build_conclusion(
        pairs=pairs,
        matches=matches,
        coverage=coverage,
        vcf_ok=vcf_ok,
        db_ok=db_ok,
        coverage_ok=coverage_ok,
    )
    replacements["_CONCLUSIONSDELANALYSE_"] = first_nonempty(
        args.conclusion, generated_conclusion
    )
    replacements.update(parse_set_values(args.set))

    replace_all(doc, replacements)
    try:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(args.output)
    except Exception as exc:
        print(f"ERROR: cannot write clinical report {args.output}: {exc}", file=sys.stderr)
        return 2

    print(f"Clinical report written to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())